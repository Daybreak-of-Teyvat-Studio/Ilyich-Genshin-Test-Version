# -*- coding: utf-8 -*-
"""
按主写的涂色图，把州装上对应的国家（owner + 核心）。

  python gamma_apply_colors.py 涂色图.png [--dry]

规则（与主写约定一致）：
  · owner 改为涂色对应的 tag；add_core_of 清掉原核心，只留新 tag
  · 逐省取主导填充色判国；州按「多数省」定国
  · 缺省跳过须弥系与纳塔系（--include-all 可取消）
  · 部分省未涂色的州、票数相同的州 → 跳过并报告，不动它
  · 天空岛（天理 PRI 的州）不动
  · 文件保持 CRLF + tab + 无 BOM；写盘前备份

产出：
  Gamma_地图/划分结果/装国结果.xlsx   每州一行：原owner / 新owner / 依据 / 处理结果
  Gamma_地图/划分结果/装国后预览.png   按新归属上色
"""
import os, re, sys, shutil, datetime, collections
import numpy as np
from PIL import Image
Image.MAX_IMAGE_PIXELS = None
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gamma_palette import COL, TAG_OF_COLOR, SUM_FACTION, NAT_FACTION, hx, dist
from gamma_regions import tag_cn, G, ROOT, OTHER, region_of, region_order

ST = os.path.join(G, 'history', 'states')
OUTDIR = os.path.join(ROOT, 'Gamma_地图', '划分结果')
os.makedirs(OUTDIR, exist_ok=True)

argv = [a for a in sys.argv[1:] if not a.startswith('--')]
DRY = '--dry' in sys.argv
ALL = '--include-all' in sys.argv
if not argv:
    print('用法: python gamma_apply_colors.py 涂色图.png [--dry] [--include-all]')
    sys.exit(1)
SRC = argv[0]

SKIP_TAGS = set() if ALL else (set(SUM_FACTION) | set(NAT_FACTION))
CONF_MIN = 0.90

# ---------------- 州 ----------------
s2o, s2p, raw = {}, {}, {}
for f in os.listdir(ST):
    t = open(os.path.join(ST, f), encoding='utf-8-sig', newline='').read()
    sid = int(re.search(r'\bid\s*=\s*(\d+)', t).group(1))
    s2o[sid] = re.search(r'\bowner\s*=\s*(\w+)', t).group(1)
    s2p[sid] = [int(x) for x in re.search(r'provinces = \{([^}]*)\}', t).group(1).split()]
    raw[sid] = t
p2s = {p: s for s, ps in s2p.items() for p in ps}
SKY = {s for s, o in s2o.items() if o == 'PRI'}

# ---------------- 省几何 ----------------
rgb2id, kind = {}, {}
for line in open(os.path.join(G, 'map', 'definition.csv'), encoding='utf-8-sig', errors='replace'):
    a = line.split(';')
    if len(a) > 6 and a[0].strip().isdigit():
        rgb2id[(int(a[1]), int(a[2]), int(a[3]))] = int(a[0])
        kind[int(a[0])] = a[4]
arr = np.asarray(Image.open(os.path.join(G, 'map', 'provinces.bmp')).convert('RGB')).astype(np.uint32)
H, W = arr.shape[:2]
key = (arr[:, :, 0] << 16) | (arr[:, :, 1] << 8) | arr[:, :, 2]
uniq, inv = np.unique(key, return_inverse=True)
plut = np.zeros(len(uniq), np.int32)
klut = np.zeros(len(uniq), np.uint8)
for i, k in enumerate(uniq):
    k = int(k)
    pid = rgb2id.get((k >> 16, (k >> 8) & 255, k & 255), 0)
    plut[i] = pid
    klut[i] = 1 if kind.get(pid, 'sea') == 'land' else 0
prov = plut[inv].reshape(H, W)
land = klut[inv].reshape(H, W).astype(bool)

# ---------------- 读涂色 ----------------
im = Image.open(SRC)
if im.size != (W, H):
    print(f'!! 涂色图尺寸 {im.size} 应为 {(W, H)}'); sys.exit(1)
rgba = np.asarray(im.convert('RGBA')).astype(np.int16)
filled = rgba[:, :, 3] > 128
ca = np.asarray(im.convert('RGB')).astype(np.int32)
col = (ca[:, :, 0] << 16) | (ca[:, :, 1] << 8) | ca[:, :, 2]

fp, fl, fc, ff = prov.ravel(), land.ravel(), col.ravel(), filled.ravel()
order = np.argsort(fp, kind='stable')
sp, sl, sf, sc = fp[order], fl[order], ff[order], fc[order]
idx = np.searchsorted(sp, np.arange(int(fp.max()) + 2))

PALV = {t: hx(c) for t, c in COL.items()}


def near_tag(c):
    k = f'#{c:06X}'
    if k in TAG_OF_COLOR:
        return TAG_OF_COLOR[k], 0.0
    rgb = ((c >> 16) & 255, (c >> 8) & 255, c & 255)
    best, bd = None, 1e9
    for t, v in PALV.items():
        d = dist(rgb, v)
        if d < bd:
            best, bd = t, d
    return (best, bd) if bd <= 60 else (None, bd)


pinfo, fuzzy = {}, {}
for pid in range(1, int(fp.max()) + 1):
    lo, hi = idx[pid], idx[pid + 1]
    if hi <= lo:
        continue
    m = sl[lo:hi]
    if not m.any():
        continue
    cols = sc[lo:hi][m]
    fl_ = sf[lo:hi][m]
    total = cols.size
    if fl_.sum() == 0:
        pinfo[pid] = (None, None, 0.0)
        continue
    cc = collections.Counter(cols[fl_].tolist())
    c, n = cc.most_common(1)[0]
    tag, d = near_tag(c)
    if d > 0:
        fuzzy[pid] = (f'#{c:06X}', tag)
    pinfo[pid] = (tag, f'#{c:06X}', n / total)
print(f'涂色图 {SRC}｜判定省 {len(pinfo)}｜非精确色板色 {len(fuzzy)} 个省')

# ---------------- 分类 ----------------
res = {}          # sid -> (结果, 新owner, 说明)
for s in sorted(s2p):
    if s in SKY:
        res[s] = ('跳过·天空岛', None, '')
        continue
    ps = s2p[s]
    tags = [pinfo[p][0] for p in ps if p in pinfo and pinfo[p][0]]
    if not tags:
        res[s] = ('跳过·未涂色', None, f'现 owner {s2o[s]}')
        continue
    cc = collections.Counter(tags)
    top, tn = cc.most_common(1)[0]
    tie = len([1 for t, n in cc.items() if n == tn]) > 1
    full = (len(tags) == len(ps))
    same = (len(cc) == 1)
    desc = '  '.join(f'{t}x{n}' for t, n in cc.most_common(4))
    if top in SKIP_TAGS:
        res[s] = ('跳过·属须弥/纳塔系', None, desc)
        continue
    if not full:
        res[s] = ('跳过·部分未涂色', None, f'已涂 {len(tags)}/{len(ps)}  {desc}')
        continue
    if tie:
        tied = [t for t, n in cc.items() if n == tn]
        keep = '原owner就是并列国之一，等于已对' if s2o[s] in tied else '原owner不在并列中，需你定'
        res[s] = ('保持原样·票数相同', None, desc + ' | ' + keep)
        continue
    res[s] = ('装上' if same else '装上·跨色按多数', top, desc)

cnt = collections.Counter(r[0] for r in res.values())
print()
for k, v in cnt.most_common():
    print(f'  {k}: {v}')

# ---------------- 应用 ----------------
def set_owner(t, tag):
    t2 = re.sub(r'\t\towner = \w+', f'\t\towner = {tag}', t, count=1)
    if t2 == t and f'\t\towner = {tag}' not in t:
        return None
    t2 = re.sub(r'\r\n\t\tadd_core_of = \w+', '', t2)
    return re.sub(r'(\t\towner = ' + tag + r')', rf'\1\r\n\t\tadd_core_of = {tag}', t2, count=1)


changed = {}
for s, (kind_, tag, desc) in res.items():
    if tag is None:
        continue
    nt = set_owner(raw[s], tag)
    if nt is None:
        print(f'  !! state {s} 未找到 owner 行，跳过')
        continue
    changed[s] = nt

# ---------------- 写盘 ----------------
if changed and not DRY:
    bdir = os.path.join(ROOT, '.backups', 'gamma_owners_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    os.makedirs(bdir, exist_ok=True)
    for s, t in changed.items():
        fn = f'{s}-State_{s}.txt'
        shutil.copy2(os.path.join(ST, fn), os.path.join(bdir, fn))
        with open(os.path.join(ST, fn), 'w', encoding='utf-8', newline='') as fh:
            fh.write(t)
    print(f'\n已写入 {len(changed)} 个州文件，备份于 {bdir}')
else:
    print(f'\n[{"DRY" if DRY else "无改动"}] 未写盘')

# ---------------- 校验 ----------------
errs = []
for s, t in changed.items():
    if t.count('{') != t.count('}'):
        errs.append(f'{s}: 花括号不配对')
    if re.search(r'(?m)^ +', t):
        errs.append(f'{s}: 行首空格')
    ows = re.findall(r'\t\towner = (\w+)', t)
    cor = re.findall(r'\t\tadd_core_of = (\w+)', t)
    if len(ows) != 1:
        errs.append(f'{s}: owner 行数 {len(ows)}')
    if cor != [res[s][1]]:
        errs.append(f'{s}: 核心 {cor} != [{res[s][1]}]')
if not DRY:
    for f in os.listdir(ST):
        rb = open(os.path.join(ST, f), 'rb').read()
        if rb.count(b'\n') != rb.count(b'\r\n'):
            errs.append(f'{f}: 行尾非纯 CRLF')
        if rb[:3] == b'\xef\xbb\xbf':
            errs.append(f'{f}: 出现 BOM')
print(f'校验：{len(errs)} 个问题')
for e in errs[:20]:
    print('  !', e)

# ---------------- 结果表 ----------------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '装国结果'
HEAD = ['state_id', '结果', '原owner', '新owner', '省数', '涂色分布', '备注']
ws.append(HEAD)
thin = Side(style='thin', color='BFBFBF'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
for c in range(1, len(HEAD) + 1):
    cell = ws.cell(1, c); cell.fill = PatternFill('solid', fgColor='305496')
    cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    cell.alignment = Alignment(horizontal='center', vertical='center'); cell.border = BD
for r, s in enumerate(sorted(res), 2):
    kind_, tag, desc = res[s]
    fz = [f'{p}:{fuzzy[p][0]}->{fuzzy[p][1]}' for p in s2p[s] if p in fuzzy]
    note = ('非色板色 ' + ' '.join(fz)) if fz else ''
    vals = [s, kind_, s2o[s], tag or '', len(s2p[s]), desc, note]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.font = Font(name='Consolas', size=10) if c in (1, 3, 4) else Font(name='微软雅黑', size=10)
        cell.alignment = Alignment(horizontal='left' if c == 6 else 'center', vertical='center')
        cell.border = BD
        if kind_.startswith('装上'):
            cell.fill = PatternFill('solid', fgColor='E2EFDA')
        elif kind_.startswith('跳过·部分') or kind_.startswith('跳过·票数'):
            cell.fill = PatternFill('solid', fgColor='FCE4D6')
for col, w in zip('ABCDEFG', (10, 20, 10, 10, 8, 46, 26)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{len(res)+1}'
wb.save(os.path.join(OUTDIR, '装国结果.xlsx'))
print(f'结果表: {os.path.join(OUTDIR, "装国结果.xlsx")}')

# ---------------- 预览图（按大区上色）----------------
REG_OF = region_of()
newown = {s: (res[s][1] or s2o[s]) for s in s2p}
rmap = {r: i + 1 for i, r in enumerate(region_order())}
lutp = np.zeros(int(prov.max()) + 1, np.int32)
for p, s in p2s.items():
    lutp[p] = rmap.get(REG_OF.get(newown[s], OTHER), rmap[OTHER])
RC = {'蒙德': (126, 211, 143), '璃月': (240, 190, 90), '稻妻': (176, 132, 224),
      '须弥': (110, 178, 96), '枫丹': (98, 168, 232), '纳塔': (232, 120, 88),
      '挪德卡莱': (231, 148, 196), '至冬': (128, 216, 224), OTHER: (160, 160, 165)}
out = np.full((H, W, 3), (240, 243, 247), np.uint8)
out[~land] = (240, 243, 247)
per = lutp[prov]
for rn, i in rmap.items():
    out[land & (per == i)] = RC.get(rn, (160, 160, 165))
Image.fromarray(out, 'RGB').save(os.path.join(OUTDIR, '装国后预览.png'))
Image.fromarray(out, 'RGB').resize((2048, 1024), Image.LANCZOS).save(
    os.path.join(OUTDIR, '装国后预览_半尺寸.png'))
print(f'预览: {os.path.join(OUTDIR, "装国后预览.png")}')
