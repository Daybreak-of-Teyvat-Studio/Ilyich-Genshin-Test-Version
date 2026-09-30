# -*- coding: utf-8 -*-
"""Beta 州名标注图（2 倍分辨率 11264x4096，字号不变 = 相对更小更清晰）"""
import os, re, sys, collections
import numpy as np
from PIL import Image, ImageDraw, ImageFont
Image.MAX_IMAGE_PIXELS = None
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
ROOT = r'C:\Users\LR\Documents\GitHub\Ilyich-Genshin-Test-Version'
BETA = os.path.join(ROOT, 'Daybreak of Teyvat Beta Version')
STB = os.path.join(BETA, 'history', 'states')

# 中文名
wb = openpyxl.load_workbook(os.path.join(ROOT, 'Beta州_国家分组表.xlsx'), read_only=True)
ws = wb.active
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
i_id, i_cn = hdr.index('state_id') + 1, hdr.index('state中文名') + 1
cn = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[i_id - 1] is not None:
        cn[int(row[i_id - 1])] = row[i_cn - 1]

s2p, s2name = {}, {}
for f in os.listdir(STB):
    t = open(os.path.join(STB, f), encoding='utf-8-sig', errors='replace').read()
    sid = int(re.search(r'\bid\s*=\s*(\d+)', t).group(1))
    s2p[sid] = [int(x) for x in re.search(r'provinces\s*=\s*\{([^}]*)\}', t).group(1).split()]
    m = re.search(r'name\s*=\s*"([^"]*)"', t)
    s2name[sid] = m.group(1) if m else ''

rgb2id, kind = {}, {}
for line in open(os.path.join(BETA, 'map', 'definition.csv'), encoding='utf-8-sig', errors='replace'):
    a = line.split(';')
    if len(a) > 6 and a[0].strip().isdigit():
        rgb2id[(int(a[1]), int(a[2]), int(a[3]))] = int(a[0])
        kind[int(a[0])] = a[4]
arr = np.asarray(Image.open(os.path.join(BETA, 'map', 'provinces.bmp')).convert('RGB')).astype(np.uint32)
H, W = arr.shape[:2]
key = (arr[:, :, 0] << 16) | (arr[:, :, 1] << 8) | arr[:, :, 2]
uniq, inv = np.unique(key, return_inverse=True)
plut = np.zeros(len(uniq), np.int32)
for i, k in enumerate(uniq):
    k = int(k)
    plut[i] = rgb2id.get((k >> 16, (k >> 8) & 255, k & 255), 0)
prov = plut[inv].reshape(H, W)
land = np.isin(prov, [p for p, k in kind.items() if k == 'land'])

flat = prov.ravel()
npx = np.bincount(flat, minlength=int(prov.max()) + 2)
gx = np.tile(np.arange(W, dtype=np.float64), H)
gy = np.repeat(np.arange(H, dtype=np.float64), W)
cx = np.bincount(flat, weights=gx, minlength=len(npx)) / np.maximum(npx, 1)
cy = np.bincount(flat, weights=gy, minlength=len(npx)) / np.maximum(npx, 1)

state = np.zeros_like(prov)
for s, ps in s2p.items():
    state[prov == ps[0] if False else np.isin(prov, ps)] = s
base = np.full((H, W, 3), (247, 248, 250), np.uint8)
base[~land] = (225, 232, 240)
sb = np.zeros((H, W), bool)
sb[:, 1:] |= (state[:, 1:] != state[:, :-1])
sb[1:, :] |= (state[1:, :] != state[:-1, :])
sb &= land
base[sb] = (70, 70, 75)
pb = np.zeros((H, W), bool)
pb[:, 1:] |= (prov[:, 1:] != prov[:, :-1])
pb[1:, :] |= (prov[1:, :] != prov[:-1, :])
pb &= land
base[pb] = (205, 205, 208)

# 2 倍 NEAREST 放大（硬边不糊），质心坐标 ×2
Z = 2
img = Image.fromarray(base, 'RGB').resize((W * Z, H * Z), Image.NEAREST)
d = ImageDraw.Draw(img)
f = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc', 13)   # 字号不变 → 相对减半
for sid in sorted(s2p):
    name = cn.get(sid, s2name[sid])
    if not name:
        continue
    xs = [cx[p] for p in s2p[sid]]
    ys = [cy[p] for p in s2p[sid]]
    tx, ty = sum(xs) / len(xs) * Z, sum(ys) / len(ys) * Z
    d.text((tx, ty), name, font=f, fill=(150, 20, 20), anchor='mm',
           stroke_width=2, stroke_fill=(255, 255, 255))
OUT = os.path.join(ROOT, 'Beta州_中文名标注图.png')
img.save(OUT)
img.resize((4096, 2048), Image.LANCZOS).save(os.path.join(ROOT, 'Beta州_中文名标注图_预览.png'))
print(f'已出: {OUT}（{W*Z}x{H*Z}）')
