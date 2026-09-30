# -*- coding: utf-8 -*-
"""
Gamma 版地图工具。产出到 Gamma_地图/ 目录：

  01_大区划分.png / _半尺寸.png     陆地按大区上色 + 州界
  02_州号总图.png / _半尺寸.png     州界 + 州号 + 坐标网格 + 省界
  03_州号_<大区>.png                每个大区放大 2 倍，州号清晰可读（含坐标网格）

并输出 Gamma州列表.xlsx（大区 / 国家名 / tag / state_id / 省数 / 质心 / 省列表，带筛选）

大区归属依据 common/scripted_triggers/DOT_scripted_triggers.txt，其中 Is_NAT 含
本次新增的 6 个纳塔部族 NCE/NMN/NPS/NFF/NSC/NCP。
"""
import os, re, sys, collections
import numpy as np
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
Image.MAX_IMAGE_PIXELS = None
sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gamma_regions import (load_regions, region_of, rank_of, tag_cn,
                           region_order, all_country_tags, G, ROOT, OTHER)

ST = os.path.join(G, 'history', 'states')
OUT = os.path.join(ROOT, 'Gamma_地图')
os.makedirs(OUT, exist_ok=True)
REGIONS = load_regions()                 # [(大区, 主国, [属国...])]，真源见 gamma_regions.py
REGION_OF = region_of()
RANK = rank_of()
REGIONS_ORDER = region_order(REGIONS)
COLOR = {
    '蒙德': (126, 211, 143), '璃月': (240, 190, 90), '稻妻': (176, 132, 224),
    '须弥': (110, 178, 96), '枫丹': (98, 168, 232), '纳塔': (232, 120, 88),
    '挪德卡莱': (231, 148, 196), '至冬': (128, 216, 224), OTHER: (160, 160, 165),
    'SEA': (240, 243, 247), 'NOPROV': (252, 250, 245),
}

# ---------------- 读州 ----------------
state2owner, state2provs = {}, {}
for f in os.listdir(ST):
    t = open(os.path.join(ST, f), encoding='utf-8-sig', errors='replace').read()
    sid = int(re.search(r'\bid\s*=\s*(\d+)', t).group(1))
    state2owner[sid] = re.search(r'\bowner\s*=\s*(\w+)', t).group(1)
    state2provs[sid] = [int(x) for x in
                        re.search(r'provinces = \{([^}]*)\}', t).group(1).split()]
prov2state = {p: s for s, ps in state2provs.items() for p in ps}
unknown = sorted({o for o in state2owner.values() if o not in REGION_OF})
print(f'州 {len(state2owner)} 个；未归类 tag（归入其他国家）: {unknown}')

# ---------------- 读 definition.csv ----------------
rgb2id, id2kind = {}, {}
for line in open(os.path.join(G, 'map', 'definition.csv'), encoding='utf-8-sig', errors='replace'):
    a = line.split(';')
    if len(a) > 6 and a[0].strip().isdigit():
        rgb2id[(int(a[1]), int(a[2]), int(a[3]))] = int(a[0])
        id2kind[int(a[0])] = a[4]

# ---------------- 像素 -> 省 -> 州 ----------------
arr = np.asarray(Image.open(os.path.join(G, 'map', 'provinces.bmp')).convert('RGB')).astype(np.uint32)
H, W = arr.shape[:2]
key = (arr[:, :, 0] << 16) | (arr[:, :, 1] << 8) | arr[:, :, 2]
uniq, inv = np.unique(key, return_inverse=True)
prov_lut = np.zeros(len(uniq), np.int32)
for i, k in enumerate(uniq):
    k = int(k)
    prov_lut[i] = rgb2id.get((k >> 16, (k >> 8) & 255, k & 255), 0)
prov = prov_lut[inv].reshape(H, W)
state = np.zeros_like(prov)
for p, s in prov2state.items():
    state[prov == p] = s

kind_lut = np.zeros(len(uniq), np.uint8)          # 1=海 2=陆
reg_lut = np.zeros(len(uniq), np.uint8)
for i in range(len(uniq)):
    pid = int(prov_lut[i])
    if pid == 0:
        continue
    if id2kind.get(pid, 'sea') == 'land':
        kind_lut[i] = 2
        own = state2owner.get(prov2state.get(pid, -1), '')
        reg_lut[i] = REGIONS_ORDER.index(REGION_OF.get(own, OTHER))
    else:
        kind_lut[i] = 1
kind = kind_lut[inv].reshape(H, W)
reg_idx = reg_lut[inv].reshape(H, W)
land = kind == 2

# 省界 / 州界
pb = np.zeros((H, W), bool)
pb[:, 1:] |= (prov[:, 1:] != prov[:, :-1])
pb[1:, :] |= (prov[1:, :] != prov[:-1, :])
sb = np.zeros((H, W), bool)
sb[:, 1:] |= (state[:, 1:] != state[:, :-1])
sb[1:, :] |= (state[1:, :] != state[:-1, :])
pb &= land
sb &= land

# 每州像素数 / 质心（bincount 一次算完）
flat = state.ravel()
gx = np.tile(np.arange(W, dtype=np.float64), H)
gy = np.repeat(np.arange(H, dtype=np.float64), W)
npix = np.bincount(flat, minlength=862)
cx = np.bincount(flat, weights=gx, minlength=862) / np.maximum(npix, 1)
cy = np.bincount(flat, weights=gy, minlength=862) / np.maximum(npix, 1)

# ---------------- 01 大区图 ----------------
base = np.full((H, W, 3), COLOR['NOPROV'], np.uint8)
base[kind == 1] = COLOR['SEA']
for i, rn in enumerate(REGIONS_ORDER):
    base[land & (reg_idx == i)] = COLOR[rn]
region_img = base.copy()
region_img[sb] = (60, 60, 60)
Image.fromarray(region_img, 'RGB').save(os.path.join(OUT, '01_大区划分.png'))
Image.fromarray(region_img, 'RGB').resize((2048, 1024), Image.LANCZOS).save(
    os.path.join(OUT, '01_大区划分_半尺寸.png'))
print('01 大区划分.png')

# ---------------- 02 州号总图 ----------------
under = (base * 0.45 + 255 * 0.55).astype(np.uint8)
under[~land] = (247, 248, 250)
under[pb] = (205, 205, 208)
under[sb] = (70, 70, 75)

FONT = r'C:\Windows\Fonts\arialbd.ttf'


def label_states(img, ids, size, x0=0, y0=0, z=1):
    """在第 ids 个州的质心处标州号；(x0,y0,z) 把原图坐标换算到本图坐标"""
    d = ImageDraw.Draw(img)
    try:
        f = ImageFont.truetype(FONT, size)
    except OSError:
        f = ImageFont.load_default()
    for sid in ids:
        x = int(round((cx[sid] - x0) * z))
        y = int(round((cy[sid] - y0) * z))
        if not (0 <= x < img.width and 0 <= y < img.height):
            continue
        s = str(sid)
        b = d.textbbox((x, y), s, font=f, anchor='mm')
        d.rectangle([b[0] - 2, b[1] - 1, b[2] + 2, b[3] + 1], fill=(255, 255, 255))
        d.text((x, y), s, font=f, fill=(190, 20, 20), anchor='mm')
    return d


full = Image.fromarray(under, 'RGB')
d = label_states(full, sorted(state2provs), 15)
for x in range(0, W, 512):
    d.line([(x, 0), (x, H)], fill=(120, 160, 220), width=1)
    d.text((x + 4, 4), str(x), font=ImageFont.truetype(FONT, 18), fill=(40, 80, 180))
for y in range(0, H, 512):
    d.line([(0, y), (W, y)], fill=(120, 160, 220), width=1)
    d.text((4, y + 4), str(y), font=ImageFont.truetype(FONT, 18), fill=(40, 80, 180))
full.save(os.path.join(OUT, '02_州号总图.png'))
full.resize((2048, 1024), Image.LANCZOS).save(os.path.join(OUT, '02_州号总图_半尺寸.png'))
print('02 州号总图.png')

# ---------------- 03 分区放大图 ----------------
for i, rn in enumerate(REGIONS_ORDER):
    m = land & (reg_idx == i)
    ys, xs = np.nonzero(m)
    if len(ys) == 0:
        continue
    pad, z = 30, 2
    x0, x1 = max(0, xs.min() - pad), min(W, xs.max() + pad + 1)
    y0, y1 = max(0, ys.min() - pad), min(H, ys.max() + pad + 1)
    crop = Image.fromarray(under[y0:y1, x0:x1], 'RGB').resize(
        ((x1 - x0) * z, (y1 - y0) * z), Image.NEAREST)
    ids = [s for s in sorted(state2provs)
           if x0 <= cx[s] < x1 and y0 <= cy[s] < y1]
    d = label_states(crop, ids, 15 * z, x0, y0, z)
    for gx0 in range(0, W, 256):                       # 坐标网格（原图坐标）
        if x0 <= gx0 < x1:
            d.line([((gx0 - x0) * z, 0), ((gx0 - x0) * z, crop.height)],
                   fill=(120, 160, 220), width=1)
            d.text(((gx0 - x0) * z + 4, 4), str(gx0),
                   font=ImageFont.truetype(FONT, 16 * z), fill=(40, 80, 180))
    for gy0 in range(0, H, 256):
        if y0 <= gy0 < y1:
            d.line([(0, (gy0 - y0) * z), (crop.width, (gy0 - y0) * z)],
                   fill=(120, 160, 220), width=1)
            d.text((4, (gy0 - y0) * z + 4), str(gy0),
                   font=ImageFont.truetype(FONT, 16 * z), fill=(40, 80, 180))
    fn = f'03_州号_{i+1:02d}_{rn}.png'
    crop.save(os.path.join(OUT, fn))
    print(f'   {fn}  {crop.width}x{crop.height}  州 {len(ids)} 个')

# ---------------- Gamma州列表.xlsx ----------------
rows = []
for sid in sorted(state2provs):
    own = state2owner[sid]
    rn = REGION_OF.get(own, OTHER)
    rows.append((rn, REGIONS_ORDER.index(rn), RANK.get(own, 99), tag_cn(own), own,
                 sid, len(state2provs[sid]), round(cx[sid]), round(cy[sid]),
                 ' '.join(map(str, sorted(state2provs[sid])))))
rows.sort(key=lambda r: (r[1], r[2], r[5]))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Gamma州'
HEAD = ['大区', '国家名', '国家tag', 'state_id', '省数', '质心x', '质心y', 'provinces']
ws.append(HEAD)
thin = Side(style='thin', color='BFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='305496')
HFONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
CEN = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
for c in range(1, 9):
    cell = ws.cell(1, c)
    cell.fill, cell.font, cell.alignment, cell.border = HF, HFONT, CEN, BD
for r, row in enumerate(rows, 2):
    for c, v in enumerate([row[0], row[3], row[4], row[5], row[6], row[7], row[8], row[9]], 1):
        cell = ws.cell(r, c, v)
        cell.font = Font(name='Consolas', size=11) if c in (2, 3) else Font(name='微软雅黑', size=11)
        cell.alignment = LEFT if c == 8 else CEN
        cell.border = BD
for col, w in zip('ABCDEFGH', (12, 18, 10, 10, 8, 9, 9, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:H{len(rows)+1}'
ws.sheet_view.showGridLines = False
wb.save(os.path.join(ROOT, 'Gamma州列表.xlsx'))
print(f'Gamma州列表.xlsx（{len(rows)} 行）')

# ---------------- 统计输出 ----------------
cnt = collections.Counter()
tagc = collections.Counter()
for sid, own in state2owner.items():
    rn = REGION_OF.get(own, OTHER)
    cnt[rn] += 1
    tagc[(rn, own)] += 1
print('\n大区 -> 国家 -> 州数')
for rn in REGIONS_ORDER:
    tags = sorted([t for (r_, t) in tagc if r_ == rn], key=lambda t: RANK.get(t, 99))
    print(f'  {rn}（{cnt[rn]} 州）: ' + '  '.join(f'{tag_cn(t)}={t} {tagc[(rn,t)]}' for t in tags))
print(f'  合计 {sum(cnt.values())} 州')
no_land = [t for t in all_country_tags() if t not in state2owner.values()]
print('\n有国家定义但 Gamma 无一州:', '  '.join(f'{tag_cn(t)}={t}' for t in no_land))
