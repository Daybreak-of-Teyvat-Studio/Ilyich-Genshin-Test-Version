# -*- coding: utf-8 -*-
"""export_migration_excel.py —— 生成 beta→gamma id 迁移对照 Excel（桌面）
Sheet: 总览 / state引用明细 / province引用明细 / 州映射表 / VP映射表 / 待人工清单"""
import os, re, sys, json, collections

sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

br = json.load(open(os.path.join(ROOT, 'tools', 'beta_gamma_bridge.json'), encoding='utf-8'))
rows = json.load(open(os.path.join(ROOT, 'tools', 'beta_refs_inventory.json'), encoding='utf-8'))
SB, PB = br['state_bridge'], br['prov_bridge']
OUT = os.path.join(os.path.expanduser('~'), 'Desktop', 'beta至gamma_id迁移对照表.xlsx')

wb = Workbook()
GREEN = PatternFill('solid', fgColor='C6EFCE')
RED = PatternFill('solid', fgColor='FFC7CE')
YELL = PatternFill('solid', fgColor='FFEB9C')
HDR = Font(bold=True)


def fill_sheet(ws, headers, data, widths, status_col=None):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).font = HDR
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in data:
        ws.append(row)
        if status_col:
            v = str(row[status_col - 1] or '')
            cell = ws.cell(ws.max_row, status_col)
            if v.startswith('✓'):
                cell.fill = GREEN
            elif v.startswith('⚠') or v.startswith('?'):
                cell.fill = YELL
            elif v.startswith('✗'):
                cell.fill = RED
    ws.freeze_panes = 'A2'


# ---- 总览 ----
ws = wb.active
ws.title = '总览'
stat = collections.Counter(r['status'].split('（')[0] for r in rows)
ws.append(['beta→gamma id 迁移对照表（生成于本轮扫描，只读未改任何文件）'])
ws['A1'].font = Font(bold=True, size=12)
info = [
    ('扫描文件数', 'gamma mod 全部 .txt 脚本（除 states/buildings/localisation）'),
    ('引用总数', len(rows)),
    ('桥规则', '州: beta州英文名→beta中文loc→gamma同名州（fallback: 同名VP→省→州）；'
               '省: beta VICTORY_POINTS 中文名 → gamma同名VP → 省号'),
    ('可直接迁移 ✓', sum(v for k, v in stat.items() if k.startswith('✓'))),
    ('被阻塞 ✗（gamma缺同名）', sum(v for k, v in stat.items() if k.startswith('✗'))),
    ('需人工 ⚠/?', sum(v for k, v in stat.items() if k.startswith('⚠') or k.startswith('?'))),
    ('迁移方案', '①先给 gamma 缺名州补名（或提供人工映射），把 ✗ 行转 ✓；'
               '②重名/多候选行由你在本表指定正确 gamma id；'
               '③确认后运行 apply_beta_to_gamma_ids.py --dry 预览 → 正式执行（自动备份、跳过注释行）'),
]
for k, v in info:
    ws.append([k, v])
for k, v in stat.most_common():
    ws.append([f'状态分布: {k}', v])
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 110

# ---- state 引用明细 ----
ws = wb.create_sheet('state引用明细')
data = [[r['file'], r['line'], r['pattern'], r['snippet'], r['id'], r['beta_name'] or '',
         r['gamma_id'] if r['gamma_id'] else '', r['status']] for r in rows if r['kind'] == '州']
fill_sheet(ws, ['文件路径', '行号', '模式', '片段', 'beta州号', 'beta中文名', 'gamma州号', '状态'],
           data, [42, 7, 17, 60, 9, 18, 10, 30], status_col=8)

# ---- province 引用明细 ----
ws = wb.create_sheet('province引用明细')
data = [[r['file'], r['line'], r['pattern'], r['snippet'], r['id'], r['beta_name'] or '',
         r['gamma_id'] if r['gamma_id'] else '', r['status']] for r in rows if r['kind'] == '省']
fill_sheet(ws, ['文件路径', '行号', '模式', '片段', 'beta省号', 'beta中文名', 'gamma省号', '状态'],
           data, [42, 7, 17, 60, 9, 18, 10, 30], status_col=8)

# ---- 州映射表 ----
ws = wb.create_sheet('州映射表')
data = [[int(k), v['en'] or '', v['cn'] or '', v['gamma_id'] if v['gamma_id'] else '',
         v['status'], v.get('note', '')] for k, v in sorted(SB.items(), key=lambda kv: int(kv[0]))]
fill_sheet(ws, ['beta州号', 'beta英文名', 'beta中文名', 'gamma州号', '状态', '备注'],
           data, [9, 26, 20, 10, 22, 30], status_col=5)

# ---- VP 映射表 ----
ws = wb.create_sheet('VP映射表')
data = [[int(k), v['cn'], v['gamma_pid'] if v['gamma_pid'] else '', v['status'], v.get('note', '')]
        for k, v in sorted(PB.items(), key=lambda kv: int(kv[0]))]
fill_sheet(ws, ['beta省号', 'beta中文名', 'gamma省号', '状态', '备注'],
           data, [9, 22, 10, 20, 30], status_col=4)

# ---- 待人工 ----
ws = wb.create_sheet('待人工清单')
manual = [r for r in rows if not r['status'].startswith('✓')]
groups = collections.defaultdict(list)
for r in manual:
    groups[(r['kind'], r['id'], r['status'])].append(r)
data = []
for (kind, i, st), rs in sorted(groups.items(), key=lambda kv: (kv[0][1], kv[0][0])):
    files = collections.Counter(x['file'] for x in rs)
    data.append([kind, i, rs[0]['beta_name'] or '', st, sum(files.values()),
                 '；'.join(f'{k}×{v}' for k, v in files.most_common(5))])
fill_sheet(ws, ['类型', 'beta id', 'beta中文名', '状态', '引用次数', '涉及文件(前5)'],
           data, [8, 9, 20, 30, 9, 90], status_col=4)

wb.save(OUT)
print(f'已生成: {OUT}')
print(f'state引用 {len([r for r in rows if r["kind"] == "州"])} 条，province引用 '
      f'{len([r for r in rows if r["kind"] == "省"])} 条，待人工分组 {len(groups)} 组')
